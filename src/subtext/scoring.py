"""Blind scoring sets, Excel sheets for the second scorer, and score files.

make_sets()      builds three sheets from archived runs:
                   practice  all valid pilot responses (calibration only, excluded from results)
                   primary   run 1 of every model on every segment (the primary outcome)
                   second    a random 20% of primary for the second scorer (seeded, reproducible)
                 response_id values come from the OS random source, so the key cannot be rebuilt
                 from the code. The key goes to scoring/private/ (git-ignored); its SHA-256 is
                 written to scoring/sets_manifest.json, which is committed at once.
export_excel()   writes one sheet as an .xlsx with drop-down inputs.
import_excel()   reads a filled .xlsx back into a score file with the rubric's record-sheet fields.
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import json
import math
import os
import random
import secrets
from pathlib import Path

from .materials import DERIVED
from .validate import parse_json

SCORING = Path("scoring")
SHEETS = SCORING / "sheets"
SCORES = SCORING / "scores"
PRIVATE = SCORING / "private"
KEY_PATH = PRIVATE / "key.json"
MANIFEST = SCORING / "sets_manifest.json"

E_CATEGORIES = ["F", "P", "N", "M", "X"]
C_CATEGORIES = ["A", "O", "X"]
SCORE_FIELDS = ["response_id", "set", "list", "code", "segment_id", "version", "category",
                "flag_plus", "valence_agree", "cue_cited", "context_cited",
                "scorer_id", "stage", "note", "scored_at"]


def now() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")


def load_answer_key() -> dict:
    with open(DERIVED / "answer_key.jsonl", encoding="utf-8") as f:
        return {r["code"]: r for r in map(json.loads, f)}


def is_l3(seg: dict) -> bool:
    return seg["level"].startswith("L3")


# ---------------------------------------------------------------- building sets

def _valid_records(phase: str, run_filter=None):
    """Yield the schema-valid attempt of each (model, run, code) cell in a phase."""
    for p in sorted((Path("runs") / phase).glob("*/run*/code*_a*.json")):
        rec = json.loads(p.read_text(encoding="utf-8"))
        if run_filter is not None and rec["run"] != run_filter:
            continue
        if rec["validation"]["schema_ok"]:
            yield p, rec


def _entry(rec: dict) -> dict:
    obj, _ = parse_json(rec["response"]["content"])
    return {"code": rec["code"], "interpretation": obj["interpretation"], "tone": obj["tone"],
            "quotes": [c["quote"] for c in obj["codes"]]}


def _new_id(used: set) -> str:
    while True:
        rid = "R" + secrets.token_hex(3).upper()
        if rid not in used:
            used.add(rid)
            return rid


def _order_by_segment(entries: list, rng) -> list:
    """Segments in random order; responses to one segment kept together, in random order."""
    by_code = {}
    for e in entries:
        by_code.setdefault(e["code"], []).append(e)
    codes = list(by_code)
    rng.shuffle(codes)
    out = []
    for c in codes:
        group = by_code[c][:]
        rng.shuffle(group)
        out.extend(group)
    return out


def _write_jsonl(path: Path, rows: list):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")


def make_sets(second_seed: int = 20261030, second_share: float = 0.20,
              allow_missing: bool = False) -> dict:
    if KEY_PATH.exists():
        raise RuntimeError(f"{KEY_PATH} already exists. Sets are made once; delete it only if "
                           "no scoring has started, and record why in docs/.")
    sysrand = random.SystemRandom()
    used, key = set(), {}
    sets = {}

    for set_name, phase, run in [("practice", "pilot", None), ("primary", "main", 1)]:
        entries = []
        cells = {}
        for p, rec in _valid_records(phase, run):
            cells.setdefault(rec["model"], set()).add(rec["code"])
            e = _entry(rec)
            e["response_id"] = _new_id(used)
            key[e["response_id"]] = {"set": set_name, "model": rec["model"], "run": rec["run"],
                                     "attempt": rec["attempt"], "code": rec["code"], "file": str(p)}
            entries.append(e)
        if set_name == "primary":
            n_codes = len(load_answer_key())
            missing = {m: n_codes - len(c) for m, c in cells.items() if len(c) < n_codes}
            if not cells or (missing and not allow_missing):
                raise RuntimeError(f"run 1 is incomplete: missing cells per model {missing or 'all'}")
        sets[set_name] = _order_by_segment(entries, sysrand)

    primary_ids = sorted(e["response_id"] for e in sets["primary"])
    k = math.ceil(second_share * len(primary_ids))
    chosen = set(random.Random(second_seed).sample(primary_ids, k))
    sets["second"] = _order_by_segment([e for e in sets["primary"] if e["response_id"] in chosen],
                                       sysrand)

    for name, rows in sets.items():
        for r in rows:
            r["set"] = name
        _write_jsonl(SHEETS / f"{name}.jsonl", rows)

    PRIVATE.mkdir(parents=True, exist_ok=True)
    key_bytes = json.dumps(key, indent=1, sort_keys=True).encode("utf-8")
    KEY_PATH.write_bytes(key_bytes)
    manifest = {
        "created_at": now(),
        "counts": {n: len(r) for n, r in sets.items()},
        "second_scorer_sample": {"seed": second_seed, "share": second_share, "k": k,
                                 "method": "random.Random(seed).sample(sorted primary response_ids, k)"},
        "key_file": str(KEY_PATH),
        "key_sha256": hashlib.sha256(key_bytes).hexdigest(),
        "sheets_sha256": {n: hashlib.sha256((SHEETS / f"{n}.jsonl").read_bytes()).hexdigest()
                          for n in sets},
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return manifest


def load_sheet(name: str) -> list:
    with open(SHEETS / f"{name}.jsonl", encoding="utf-8") as f:
        return [json.loads(l) for l in f]


# ---------------------------------------------------------------- score files

def valence_agree(entry: dict, seg: dict) -> int:
    ref = seg["intended_valence"] if seg["version"] == "E" else seg["surface_valence"]
    return int(entry["tone"].strip().lower() == ref.strip().lower())


def score_path(set_name: str, scorer: str) -> Path:
    return SCORES / f"{set_name}__{scorer}.csv"


def read_scores(set_name: str, scorer: str) -> dict:
    p = score_path(set_name, scorer)
    if not p.exists():
        return {}
    with open(p, encoding="utf-8", newline="") as f:
        return {r["response_id"]: r for r in csv.DictReader(f)}


def write_scores(set_name: str, scorer: str, rows: dict):
    p = score_path(set_name, scorer)
    p.parent.mkdir(parents=True, exist_ok=True)
    tmp = p.with_suffix(".tmp")
    with open(tmp, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=SCORE_FIELDS)
        w.writeheader()
        for r in rows.values():
            w.writerow({k: r.get(k, "") for k in SCORE_FIELDS})
    os.replace(tmp, p)


def make_score_row(entry: dict, seg: dict, scorer: str, category: str, flag_plus, cue_cited,
                   context_cited, note: str, stage: str = "pre-discussion") -> dict:
    """Build one record-sheet row and enforce the rubric's field rules."""
    version = seg["version"]
    allowed = E_CATEGORIES if version == "E" else C_CATEGORIES
    if category not in allowed:
        raise ValueError(f"{entry['response_id']}: category {category!r} not allowed for version {version}")
    if category == "X" and not note.strip():
        raise ValueError(f"{entry['response_id']}: X requires a reason in the note")
    blank = ""
    is_e = version == "E"
    return {
        "response_id": entry["response_id"], "set": entry["set"], "list": seg["list"],
        "code": seg["code"], "segment_id": seg["id"], "version": version, "category": category,
        "flag_plus": (int(bool(flag_plus)) if category in ("F", "P") else blank),
        "valence_agree": valence_agree(entry, seg),
        "cue_cited": (int(cue_cited) if is_e and cue_cited not in (None, "") else blank),
        "context_cited": (int(context_cited) if is_e and is_l3(seg) and context_cited not in (None, "")
                          else blank),
        "scorer_id": scorer, "stage": stage, "note": note.strip(), "scored_at": now(),
    }


# ---------------------------------------------------------------- Excel for the second scorer

INPUT_HEADERS = ["category", "flag_plus", "cue_cited", "context_cited", "note"]
SHEET_HEADERS = ["response_id", "segment_id", "version", "layer", "level", "interviewee_details",
                 "passage", "surface_meaning", "intended_meaning", "key_cues", "surface_valence",
                 "intended_valence", "author_confidence", "answer_key_notes",
                 "response_interpretation", "response_tone", "response_quotes"] + INPUT_HEADERS


def export_excel(set_name: str, out_path: str | None = None) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    rows = load_sheet(set_name)
    key = load_answer_key()
    out = Path(out_path or SCORING / "excel" / f"{set_name}_scoring_sheet.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)

    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    yellow = PatternFill("solid", start_color="FFF2CC")
    grey = PatternFill("solid", start_color="E7E6E6")
    header_fill = PatternFill("solid", start_color="D9E1F2")

    wb = Workbook()
    info = wb.active
    info.title = "Instructions"
    lines = [
        ("Scoring sheet", f"set: {set_name}   responses: {len(rows)}"),
        ("What to fill in", "Only the yellow cells in the 'Scoring' tab (columns category to note). "
                            "Grey cells are not applicable. Do not edit any other cell."),
        ("Reference", "Coding Scoring Rubric (scoring_rubric.docx), Sections 2–7, and the Stage 1 "
                      "analysis plan, Section 5.1 (implicit mismatch on E segments)."),
        ("category", "E segments: F, P, N, M or X.  C segments: A, O or X.  X needs a reason in note."),
        ("flag_plus", "1 or 0. Only for F or P (rubric Section 2, '+ over-extension flag')."),
        ("cue_cited", "1 or 0. E segments only: the response cites at least one key_cues item or an "
                      "equivalent detail (rubric 7.2)."),
        ("context_cited", "1 or 0. L3 E segments only: the response cites the preceding context."),
        ("note", "Key reasons for the decision; required for X."),
        ("Not needed", "Valence agreement is computed automatically from response_tone."),
        ("Order", "Rows are grouped by segment. Please finish all responses to one segment before "
                  "moving on (rubric Section 9)."),
        ("Independence", "Please score without seeing anyone else's scores."),
        ("Example row (format only, from rubric Section 6)",
         "B-L1-02-E | response: 'He says it's nothing, but he is very proud of and happy about the "
         "letter.' | category F | flag_plus 0 | cue_cited 1 | context_cited (blank, not L3) | "
         "note: pride + the letter, both components match"),
    ]
    for i, (a, b) in enumerate(lines, 1):
        info.cell(row=i, column=1, value=a).font = bold
        c = info.cell(row=i, column=2, value=b)
        c.font, c.alignment = font, wrap
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 110

    ws = wb.create_sheet("Scoring")
    for j, h in enumerate(SHEET_HEADERS, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill, c.alignment = bold, (yellow if h in INPUT_HEADERS else header_fill), wrap
    widths = {"response_id": 11, "segment_id": 13, "version": 8, "layer": 14, "level": 13,
              "interviewee_details": 28, "passage": 80, "surface_meaning": 32, "intended_meaning": 36,
              "key_cues": 40, "surface_valence": 11, "intended_valence": 12, "author_confidence": 11,
              "answer_key_notes": 28, "response_interpretation": 50, "response_tone": 10,
              "response_quotes": 45, "category": 10, "flag_plus": 9, "cue_cited": 9,
              "context_cited": 10, "note": 40}
    from openpyxl.utils import get_column_letter
    col = {h: get_column_letter(j) for j, h in enumerate(SHEET_HEADERS, 1)}
    for h, w in widths.items():
        ws.column_dimensions[col[h]].width = w

    dv_e = DataValidation(type="list", formula1='"F,P,N,M,X"', allow_blank=True)
    dv_c = DataValidation(type="list", formula1='"A,O,X"', allow_blank=True)
    dv_01 = DataValidation(type="list", formula1='"1,0"', allow_blank=True)
    for dv in (dv_e, dv_c, dv_01):
        ws.add_data_validation(dv)

    for i, e in enumerate(rows, 2):
        seg = key[e["code"]]
        values = {
            "response_id": e["response_id"], "segment_id": seg["id"], "version": seg["version"],
            "layer": seg["layer"], "level": seg["level"], "interviewee_details": seg["details"],
            "passage": seg["passage"], "surface_meaning": seg["surface_meaning"],
            "intended_meaning": seg["intended_meaning"],
            "key_cues": "\n".join(f"• {c}" for c in seg["key_cues"]),
            "surface_valence": seg["surface_valence"], "intended_valence": seg["intended_valence"],
            "author_confidence": seg["author_confidence"], "answer_key_notes": seg["notes"],
            "response_interpretation": e["interpretation"], "response_tone": e["tone"],
            "response_quotes": "\n".join(f"“{q}”" for q in e["quotes"]),
        }
        for h, v in values.items():
            c = ws[f"{col[h]}{i}"]
            c.value, c.font, c.alignment = v, font, wrap
        is_e = seg["version"] == "E"
        applicable = {"category": True, "flag_plus": is_e, "cue_cited": is_e,
                      "context_cited": is_e and is_l3(seg), "note": True}
        for h in INPUT_HEADERS:
            c = ws[f"{col[h]}{i}"]
            c.font, c.alignment = font, wrap
            c.fill = yellow if applicable[h] else grey
        (dv_e if is_e else dv_c).add(f"{col['category']}{i}")
        for h in ("flag_plus", "cue_cited", "context_cited"):
            if applicable[h]:
                dv_01.add(f"{col[h]}{i}")
        ws.row_dimensions[i].height = 400
    ws.freeze_panes = "C2"
    wb.save(out)
    return str(out)


def import_excel(path: str, set_name: str, scorer: str) -> dict:
    from openpyxl import load_workbook

    ws = load_workbook(path, data_only=True)["Scoring"]
    headers = [c.value for c in ws[1]]
    sheet = {e["response_id"]: e for e in load_sheet(set_name)}
    key = load_answer_key()
    out, problems = {}, []
    for row in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(headers, row))
        rid = r.get("response_id")
        if not rid:
            continue
        if rid not in sheet:
            problems.append(f"{rid}: not in set {set_name}"); continue
        entry = sheet[rid]
        cat = str(r.get("category") or "").strip().upper()
        if not cat:
            problems.append(f"{rid}: category empty"); continue

        def bit(v):
            return "" if v in (None, "") else int(str(v).strip().split(".")[0])
        try:
            out[rid] = make_score_row(entry, key[entry["code"]], scorer, cat, bit(r.get("flag_plus")),
                                      bit(r.get("cue_cited")), bit(r.get("context_cited")),
                                      str(r.get("note") or ""))
        except ValueError as ex:
            problems.append(str(ex))
    missing = sorted(set(sheet) - set(out))
    if problems or missing:
        return {"written": False, "problems": problems, "unscored": missing}
    write_scores(set_name, scorer, out)
    return {"written": True, "file": str(score_path(set_name, scorer)), "n": len(out)}
